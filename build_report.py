from pathlib import Path
from copy import deepcopy

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MEDIA = ROOT / "tmp" / "source_media" / "word" / "media"
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "012012301304-Net3_Group_11-UDM_10.docx"

BLACK = "000000"
NAVY = "1F4E78"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EDF3F8"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"

IMAGE_ALT = {
    "image1.png": "Logo Trường Đại học Giao thông vận tải Thành phố Hồ Chí Minh",
    "image2.jpeg": "Sơ đồ kiến trúc Client Server và các lớp chính của hệ thống UDM10",
    "image3.jpeg": "Sơ đồ tuần tự luồng Request Ready File Data Completed",
    "image4.png": "Giao diện Client khi nhiều file đang được xử lý",
    "image5.png": "Giao diện Client sau khi có file hoàn tất và file lỗi",
    "image6.jpeg": "Ba trạng thái giao diện Client trong quá trình upload",
    "image7.jpeg": "Các log Server ghi nhận ca thành công, lỗi và timeout",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Times New Roman", size=13, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, before=0, after=0, first_line=0.75):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if line == 1.5 else WD_LINE_SPACING.SINGLE
    if line != 1.5:
        fmt.line_spacing = line
    fmt.first_line_indent = Cm(first_line) if first_line else None


def add_body(doc, text, *, bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, line=1.5, after=0, first_line=0.75):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, align=align, line=line, after=after, first_line=first_line)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=size, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=size)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size)
    return p


def add_bullet(doc, text, level=0, size=12.5, line=1.25, after=1):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.75 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.first_line_indent = None
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 2)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 13, bold=True, italic=(level == 3))
    if level == 1:
        r.text = text.upper()
    return p


def add_caption(doc, text, kind="figure"):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = kind == "table"
    p.paragraph_format.keep_together = True
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(6 if kind == "figure" else 0)
    p.paragraph_format.space_after = Pt(3 if kind == "figure" else 6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == "figure" else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, size=13, bold=True)
    return p


def add_picture(doc, path, width, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    shape = p.add_run().add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", IMAGE_ALT.get(Path(path).name, Path(path).stem))
    return p


def add_page_break(doc):
    # Attach the break to the preceding paragraph. A standalone break paragraph
    # can be pushed onto a new page when the current page is full, creating an
    # unintended blank page before the next section.
    p = doc.paragraphs[-1]
    p.add_run().add_break(WD_BREAK.PAGE)


def set_cell_text(cell, text, *, size=10.5, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths_cm, *, font_size=10.5, header_fill=NAVY, repeat_header=True, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if repeat_header:
        set_repeat_table_header(table.rows[0])
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.width = Cm(widths_cm[j])
        set_cell_shading(cell, header_fill)
        set_cell_text(cell, h, size=font_size, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].width = Cm(widths_cm[j])
            if i % 2 == 1:
                set_cell_shading(cells[j], PALE_BLUE)
            alignment = alignments[j] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[j], value, size=font_size, align=alignment)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "26")
    r_pr.extend([fonts, color, underline, size])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    set_run_font(run, size=13)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def set_page_number_start(section, value=1):
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(value))


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(3)
section.right_margin = Cm(2)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.header_distance = Cm(1.2)
section.footer_distance = Cm(1.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(13)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.color.rgb = RGBColor(0, 0, 0)

# Some Word installations add a theme border to the built-in Title style.
# The supplied formatting guide uses spacing instead of a decorative rule.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

# Page 1 - cover
for text, size, bold, after in [
    ("BỘ XÂY DỰNG", 14, True, 0),
    ("TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI THÀNH PHỐ HỒ CHÍ MINH", 14, True, 0),
    ("KHOA CÔNG NGHỆ THÔNG TIN", 14, True, 12),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
add_picture(doc, MEDIA / "image1.png", Inches(3.7))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(12)
r = p.add_run("BÁO CÁO TIỂU LUẬN CUỐI KỲ")
set_run_font(r, size=16, bold=True)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(22)
r = p.add_run("ỨNG DỤNG UPLOAD NHIỀU FILE BẰNG KÉO THẢ QUA TCP SOCKET")
set_run_font(r, size=20, bold=True)
cover_rows = [
    ("Giảng viên hướng dẫn", "Mai Ngọc Châu"),
    ("Môn học", "Lập trình mạng"),
    ("Mã học phần", "012012301304"),
    ("Mã nhóm", "Net3_Group_11"),
    ("Mã project", "UDM_10"),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
remove_table_borders(t)
for a, b in cover_rows:
    cells = t.add_row().cells
    cells[0].width = Cm(5.0)
    cells[1].width = Cm(8.0)
    set_cell_text(cells[0], a, size=13, bold=True)
    set_cell_text(cells[1], b, size=13)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(24)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Thành phố Hồ Chí Minh, tháng 9 năm 2026")
set_run_font(r, size=13, bold=True)
add_page_break(doc)

# Page 2 - assignment
add_heading(doc, "BẢNG PHÂN CÔNG VÀ MỨC ĐỘ HOÀN THÀNH", 1)
assignment_rows = [
    ("1", "Lê Văn Nhựt", "095205005482", "WPF; chọn và kéo-thả file; trạng thái; thao tác GUI; báo cáo; slide và video demo", "100%"),
    ("2", "Phạm Anh Tuấn", "075205019210", "Shared protocol; framing; kiểm tra metadata; tài liệu kỹ thuật", "100%"),
    ("3", "Nguyễn Tấn Hiệp", "087205010642", "Client TCP; hàng đợi; giới hạn đồng thời; Cancel/Retry; tích hợp và kiểm thử lỗi", "100%"),
    ("4", "Huỳnh Anh Kiệt", "051206006174", "TCP Server; session; timeout; graceful shutdown; logging", "100%"),
    ("5", "Võ Nhật Linh", "045205006605", "Truyền theo chunk; lưu trữ; SHA-256; thống kê và performance test", "100%"),
]
add_table(doc, ["STT", "Họ tên", "MSSV", "Công việc phụ trách", "Hoàn thành"], assignment_rows,
          [1.0, 3.2, 3.2, 8.0, 2.0], font_size=10.5,
          alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                      WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])
add_heading(doc, "NGUYÊN TẮC PHỐI HỢP", 2)
add_body(doc, "Mỗi thành viên chịu trách nhiệm chính với hạng mục được phân công, đồng thời tham gia rà soát tích hợp trên cùng solution. Các thay đổi được theo dõi qua lịch sử commit của repository; kết quả cuối cùng được kiểm tra theo luồng Client - Shared - Server - Storage.")
add_body(doc, "Tỷ lệ hoàn thành trong bảng phản ánh phần việc đã bàn giao vào phiên bản hiện tại. Các hạng mục còn giới hạn hoặc chưa triển khai được nêu riêng tại mục 11 để tránh đồng nhất mức hoàn thành công việc cá nhân với mức hoàn thiện toàn bộ sản phẩm.")
add_page_break(doc)

# Page 3 - manual contents and lists
add_heading(doc, "MỤC LỤC", 1)
toc_items = [
    ("1. Lý do chọn đề tài, mục tiêu và phạm vi", "1"),
    ("2. Lý thuyết tóm tắt", "2"),
    ("3. Kiến trúc hệ thống", "3"),
    ("4. Protocol và luồng truyền file", "4"),
    ("5. Thiết kế giao diện", "5"),
    ("6. Hướng dẫn chạy chương trình", "6"),
    ("7. Kiểm thử chức năng và lỗi", "7"),
    ("8. Stress test và performance test", "8"),
    ("9. Kết quả đạt được", "9"),
    ("10. Khó khăn và kiến thức học được", "10"),
    ("11. Nội dung đã và chưa hoàn thành", "11"),
    ("12. Kết luận, liên kết và tài liệu tham khảo", "12"),
]
t = doc.add_table(rows=0, cols=2)
t.autofit = False
t.alignment = WD_TABLE_ALIGNMENT.CENTER
remove_table_borders(t)
for title, page in toc_items:
    cells = t.add_row().cells
    cells[0].width = Cm(15.5)
    cells[1].width = Cm(1.3)
    set_cell_text(cells[0], title, size=12.5)
    set_cell_text(cells[1], page, size=12.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
add_heading(doc, "DANH MỤC HÌNH", 2)
for text in [
    "Hình 3.1. Sơ đồ kiến trúc Client Server và các lớp chính ........................................ 3",
    "Hình 4.1. Sequence diagram luồng Request Ready File Data Completed .................. 4",
    "Hình 5.1. Giao diện Client trong quá trình upload .................................................. 5",
    "Hình 7.1. Log Server của ca upload thành công ..................................................... 7",
]:
    add_body(doc, text, size=11.5, line=1.0, first_line=0, after=0)
add_heading(doc, "DANH MỤC BẢNG", 2)
for text in [
    "Bảng 2.1. Cơ chế kỹ thuật chính .................................................................................. 2",
    "Bảng 3.1. Thành phần và trách nhiệm ........................................................................... 3",
    "Bảng 4.1. Trường dữ liệu của UploadRequest .............................................................. 4",
    "Bảng 6.1. Cấu hình mặc định ......................................................................................... 6",
    "Bảng 7.1. Nhóm ca kiểm thử ........................................................................................... 7",
    "Bảng 8.1. Kết quả stress test .......................................................................................... 8",
    "Bảng 11.1. Phạm vi triển khai ...................................................................................... 11",
]:
    add_body(doc, text, size=11.5, line=1.0, first_line=0, after=0)

# Start numbered main section on physical page 4
main_section = doc.add_section(WD_SECTION.NEW_PAGE)
main_section.page_width = Cm(21)
main_section.page_height = Cm(29.7)
main_section.left_margin = Cm(3)
main_section.right_margin = Cm(2)
main_section.top_margin = Cm(2.5)
main_section.bottom_margin = Cm(2.5)
main_section.header_distance = Cm(1.0)
main_section.footer_distance = Cm(1.2)
main_section.header.is_linked_to_previous = False
main_section.footer.is_linked_to_previous = False
header = main_section.header
header.paragraphs[0].text = ""
add_page_number(header.paragraphs[0])
set_page_number_start(main_section, 1)

# Page 4 / numbered page 1
add_heading(doc, "1. LÝ DO CHỌN ĐỀ TÀI MỤC TIÊU VÀ PHẠM VI", 1)
add_heading(doc, "1.1. Lý do chọn đề tài", 2)
add_body(doc, "Các bài tập socket cơ bản thường truyền chuỗi ngắn nên chưa thể hiện đầy đủ vấn đề khi truyền file thật. Bài toán upload nhiều file buộc nhóm xác định ranh giới metadata và dữ liệu, xử lý kết nối bị ngắt, dọn file dang dở, theo dõi tiến độ từng file và giữ giao diện phản hồi khi nhiều tác vụ chạy đồng thời.")
add_body(doc, "Thao tác kéo-thả phù hợp với ứng dụng desktop và tạo phạm vi vừa đủ cho môn Lập trình mạng: hệ thống có Client, Server, protocol riêng, xử lý bất đồng bộ, kiểm tra toàn vẹn và các tình huống lỗi có thể tái hiện.")
add_heading(doc, "1.2. Mục tiêu", 2)
add_body(doc, "Mục tiêu tổng quát là xây dựng ứng dụng WPF trên Windows cho phép chọn hoặc kéo-thả nhiều file và upload đến TCP Server. Mỗi file có trạng thái độc lập; lỗi hoặc thao tác hủy của một file không làm dừng toàn bộ hàng đợi.")
for item in [
    "Xây dựng protocol có framing rõ ràng, kiểm tra phiên bản, requestId, tên file, kích thước và SHA-256.",
    "Giới hạn tối đa ba upload đồng thời trên mỗi Client; hỗ trợ Cancel và Retry theo từng file.",
    "Server nhận đúng số byte, ghi file tạm, kiểm tra hash, đổi tên khi trùng và không ghi đè file cũ.",
    "Ghi nhận kết quả kiểm thử, throughput, độ trễ, CPU, RAM và tỷ lệ lỗi theo điều kiện đo cụ thể.",
]:
    add_bullet(doc, item)
add_heading(doc, "1.3. Đối tượng phạm vi và phương pháp", 2)
add_body(doc, "Đối tượng nghiên cứu gồm TCP Client, TCP Server, protocol truyền file, hàng đợi upload và cơ chế lưu trữ cục bộ. Sản phẩm tập trung vào mạng TCP trong phạm vi học phần, chưa triển khai xác thực, TLS, cloud storage hoặc resume theo offset. Nhóm phân tích yêu cầu, thiết kế protocol, lập trình theo module, kiểm thử chức năng và đo hiệu năng trên loopback.")
add_page_break(doc)

# Page 5 / numbered page 2
add_heading(doc, "2. LÝ THUYẾT TÓM TẮT", 1)
add_heading(doc, "2.1. TCP Socket và framing", 2)
add_body(doc, "TCP cung cấp luồng byte tin cậy và có thứ tự, nhưng không giữ ranh giới message cho ứng dụng [1]. Vì vậy, một lần gửi có thể được đọc thành nhiều phần hoặc nhiều phần gửi có thể xuất hiện trong cùng một lần đọc. Hệ thống dùng length prefix 4 byte little-endian cho JSON metadata và dùng fileSize để xác định chính xác phần dữ liệu nhị phân.")
add_heading(doc, "2.2. Bất đồng bộ hàng đợi và hủy tác vụ", 2)
add_body(doc, "Các thao tác kết nối, đọc, ghi và truyền dữ liệu dùng async/await để không chặn luồng giao diện. UploadManager dùng SemaphoreSlim giới hạn số tác vụ đồng thời. Mỗi file có CancellationToken riêng; thao tác Cancel đóng vòng đời truyền của file đó, còn các mục khác tiếp tục chạy.")
add_heading(doc, "2.3. Kiểm tra toàn vẹn và lưu file an toàn", 2)
add_body(doc, "SHA-256 tạo giá trị băm để phát hiện nội dung thay đổi [3]. Cơ chế này không mã hóa dữ liệu và không thay thế xác thực hoặc TLS. Server ghi dữ liệu vào file .part, chỉ chuyển thành file chính thức sau khi nhận đủ byte và hash khớp. Nếu thiếu byte, timeout hoặc checksum sai, file tạm được xóa khi tiến trình còn khả năng xử lý lỗi.")
add_caption(doc, "Bảng 2.1. Cơ chế kỹ thuật chính", kind="table")
theory_rows = [
    ("Ranh giới message", "Length prefix cho JSON; fileSize cho raw binary", "Không phụ thuộc số lần ReadAsync"),
    ("Giao diện phản hồi", "async/await và cập nhật trạng thái", "Không chặn UI thread"),
    ("Giới hạn tải", "SemaphoreSlim", "Tối đa 3 upload trên mỗi Client"),
    ("Hủy từng file", "CancellationToken và đóng socket", "Không dừng toàn bộ hàng đợi"),
    ("Toàn vẹn", "SHA-256", "Phát hiện dữ liệu nhận khác metadata"),
    ("Lưu an toàn", "File .part rồi đổi tên", "Không công nhận file chưa hoàn tất"),
]
add_table(doc, ["Vấn đề", "Cơ chế", "Ý nghĩa"], theory_rows, [4.0, 6.5, 7.0], font_size=10.2)
add_page_break(doc)

# Page 6 / numbered page 3
add_heading(doc, "3. KIẾN TRÚC HỆ THỐNG", 1)
add_body(doc, "Solution gồm ba project. UDM10.Client là ứng dụng WPF tiếp nhận thao tác người dùng và quản lý hàng đợi. UDM10.Server lắng nghe TCP, tạo session cho từng kết nối và chuyển dữ liệu đến lớp lưu trữ. UDM10.Shared chứa model, enum trạng thái, framing, validation và hàm truyền theo chunk để hai phía dùng cùng quy tắc.")
add_picture(doc, MEDIA / "image2.jpeg", Inches(6.35))
add_caption(doc, "Hình 3.1. Sơ đồ kiến trúc Client Server và các lớp chính", kind="figure")
add_caption(doc, "Bảng 3.1. Thành phần và trách nhiệm", kind="table")
architecture_rows = [
    ("UDM10.Client", "MainWindow, ViewModel, UploadQueueService, UploadManager, ClientTransfer", "Kéo-thả, lập hàng đợi, hiển thị tiến độ, Cancel/Retry và truyền file"),
    ("UDM10.Shared", "Request/Response, status, framing, validation, serializer", "Dùng chung cấu trúc message và quy tắc kiểm tra"),
    ("UDM10.Server", "TcpListener, ClientConnectionHandler, FileStorageService", "Quản lý session, timeout, lưu .part, kiểm tra hash và ghi log"),
]
add_table(doc, ["Project", "Thành phần chính", "Trách nhiệm"], architecture_rows, [3.2, 6.3, 8.0], font_size=9.7)
add_page_break(doc)

# Page 7 / numbered page 4
add_heading(doc, "4. PROTOCOL VÀ LUỒNG TRUYỀN FILE", 1)
add_body(doc, "Phiên bản protocol hiện tại là V3, cổng mặc định 9000. Metadata được mã hóa JSON UTF-8, có độ dài từ 1 đến 4096 byte. Client chỉ gửi dữ liệu file sau khi nhận response Ready đúng protocolVersion và requestId. Khi hoàn tất, Server trả Completed kèm savedFileName.")
add_caption(doc, "Bảng 4.1. Trường dữ liệu của UploadRequest", kind="table")
request_rows = [
    ("protocolVersion", "Bắt buộc bằng V3"),
    ("requestId", "ID mới cho mỗi lượt upload; dài 1-128 ký tự"),
    ("fileName", "Chỉ chứa tên file; dài 1-255 ký tự; không nhận đường dẫn hoặc tên thiết bị Windows"),
    ("fileSize", "Số byte cần nhận; không âm và không vượt giới hạn Server"),
    ("fileHash", "SHA-256 gồm 64 ký tự hex"),
    ("status", "Luôn là Request, kể cả khi Retry"),
]
add_table(doc, ["Trường", "Quy tắc"], request_rows, [4.3, 13.2], font_size=10.1)
add_picture(doc, MEDIA / "image3.jpeg", Inches(6.0))
add_caption(doc, "Hình 4.1. Sequence diagram luồng Request Ready File Data Completed", kind="figure")
add_body(doc, "Mỗi file dùng một requestId và một TCP connection. Retry tạo request mới và gửi lại từ đầu. Phiên bản hiện tại chưa hỗ trợ tiếp tục từ offset giữa file.", size=11.5, line=1.15, first_line=0.75)
add_page_break(doc)

# Page 8 / numbered page 5
add_heading(doc, "5. THIẾT KẾ GIAO DIỆN", 1)
add_body(doc, "MainWindow gồm vùng nhập IP và port, trạng thái kết nối, vùng kéo-thả, nút chọn file, thao tác hàng loạt và danh sách file. Mỗi dòng hiển thị tên nguồn, tên Server đã lưu, kích thước, trạng thái, phần trăm, tốc độ, thông báo và nút Cancel hoặc Retry.")
add_body(doc, "Trạng thái được thể hiện bằng chữ gồm Đang chờ, Đang tải, Hoàn tất, Lỗi và Đã hủy. Nút thao tác chỉ được bật khi trạng thái tương ứng hợp lệ. Dòng thống kê tổng hợp số file hoàn tất, lỗi, đã hủy, tổng byte, thời gian và tốc độ trung bình.")
gui_table = doc.add_table(rows=1, cols=3)
gui_table.alignment = WD_TABLE_ALIGNMENT.CENTER
gui_table.autofit = False
remove_table_borders(gui_table)
for cell, image_path in zip(gui_table.rows[0].cells, [MEDIA / "image4.png", MEDIA / "image5.png", MEDIA / "image6.jpeg"]):
    cell.width = Cm(5.8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    shape = p.add_run().add_picture(str(image_path), width=Inches(1.95))
    shape._inline.docPr.set("descr", IMAGE_ALT[image_path.name])
add_caption(doc, "Hình 5.1. Các trạng thái giao diện Client trong quá trình upload", kind="figure")
add_page_break(doc)

# Page 9 / numbered page 6
add_heading(doc, "6. HƯỚNG DẪN CHẠY CHƯƠNG TRÌNH", 1)
add_heading(doc, "6.1. Môi trường", 2)
add_body(doc, "Client WPF cần Windows 10/11 hoặc Windows VM có .NET 10 SDK. Server và các công cụ kiểm thử protocol có thể chạy trên nền tảng hỗ trợ .NET. Phần stress test trong báo cáo được đo trên macOS bằng Client kiểm thử; phần GUI được chạy và minh họa trong môi trường Windows.")
add_caption(doc, "Bảng 6.1. Cấu hình mặc định", kind="table")
config_rows = [
    ("Server IP", "127.0.0.1", "0.0.0.0"),
    ("Port", "9000", "9000"),
    ("Connect timeout", "5000 ms", "Không áp dụng"),
    ("Receive timeout", "30000 ms", "30000 ms"),
    ("Chunk size", "65536 byte", "65536 byte"),
    ("Số file đồng thời", "3", "Chưa giới hạn tổng"),
    ("Kích thước file tối đa", "Theo Server", "10 GiB mỗi file"),
    ("Thư mục lưu", "Không áp dụng", "Uploads"),
]
add_table(doc, ["Tham số", "Client", "Server"], config_rows, [6.7, 5.2, 5.6], font_size=9.8)
add_heading(doc, "6.2. Các bước chạy", 2)
steps = [
    "Mở PowerShell tại thư mục Code, chạy dotnet restore .\\UDM10.sln.",
    "Chạy dotnet build .\\UDM10.sln -c Release --no-restore và xác nhận không có lỗi.",
    "Chạy Server bằng dotnet run --project .\\Server\\UDM10.Server.csproj -c Release.",
    "Chạy Client bằng dotnet run --project .\\Client\\UDM10.Client.csproj -c Release.",
    "Nhập IP và port; chọn hoặc kéo-thả file. Kiểm tra kết quả trong thư mục Uploads của Server.",
]
for i, item in enumerate(steps, 1):
    add_body(doc, f"{i}. {item}", size=11.5, line=1.1, first_line=0, after=1)
add_page_break(doc)

# Page 10 / numbered page 7
add_heading(doc, "7. KIỂM THỬ CHỨC NĂNG VÀ LỖI", 1)
add_body(doc, "Các ca kiểm thử ghi rõ môi trường, commit, dữ liệu đầu vào, mức đồng thời, kết quả mong đợi, kết quả thực tế và bằng chứng. Bộ test kỹ thuật trong README ghi nhận build Release sạch với 0 warning, 0 error và đạt 11/11 ca protocol, storage, send timeout và scheduler ngày 12/09/2026.")
add_caption(doc, "Bảng 7.1. Nhóm ca kiểm thử", kind="table")
test_rows = [
    ("Chức năng", "Chọn/kéo 1, 3, 5, 20 file", "Đủ file; tối đa 3 file chạy; GUI phản hồi; có tiến độ"),
    ("File đặc biệt", "File rỗng và trùng tên", "File rỗng hợp lệ; file trùng được đổi tên; không ghi đè"),
    ("Điều khiển", "Cancel/Retry từng file và toàn bộ", "Chỉ tác động đúng mục; Retry gửi lại từ đầu"),
    ("Dữ liệu sai", "JSON, UTF-8, length, version, status, hash", "Server trả Error hoặc đóng phiên an toàn; không crash"),
    ("Ngắt kết nối", "Tắt Client hoặc dừng Server giữa upload", "Không công nhận file thiếu; dọn .part khi còn xử lý được"),
    ("Lưu trữ", "Không có quyền ghi hoặc hết dung lượng", "Trả StorageError nếu có thể; không tạo file hoàn tất giả"),
]
add_table(doc, ["Nhóm", "Ca kiểm thử", "Kết quả cần xác nhận"], test_rows, [3.0, 6.4, 8.1], font_size=9.5)
add_picture(doc, MEDIA / "image7.jpeg", Inches(2.55))
add_caption(doc, "Hình 7.1. Log Server của các ca upload thành công, lỗi và timeout", kind="figure")
add_page_break(doc)

# Page 11 / numbered page 8
add_heading(doc, "8. STRESS TEST VÀ PERFORMANCE TEST", 1)
add_body(doc, "Nhóm đo hai mức tải trên TCP loopback. Throughput bằng tổng byte truyền thành công chia cho thời gian đo; CPU và RAM được lấy trong cùng khoảng chạy. Mức tải 2 sử dụng ba Client, mỗi Client tối đa ba upload đồng thời.")
add_caption(doc, "Bảng 8.1. Kết quả stress test", kind="table")
stress_rows = [
    ("Kịch bản", "10 file x 10 MiB", "30 file x 20 MiB"),
    ("Client và concurrency", "1 Client; tối đa 3", "3 Client; tối đa 9"),
    ("Mạng và máy", "Loopback; MacBook Air M4 10 nhân; RAM 16 GB; macOS 15.7.3; .NET 10.0.400", "Giống mức tải 1"),
    ("Tổng thời gian", "0,115 giây", "0,315 giây"),
    ("Throughput", "867,44 MiB/s", "1.903,11 MiB/s"),
    ("Phản hồi ban đầu", "TB 8,13 ms; max 27,56 ms", "TB 21,77 ms; P95 66,40 ms; max 66,53 ms"),
    ("CPU/RAM Client", "CPU TB 50,90%; RAM đỉnh 17,17 MiB", "Tổng CPU TB 119,84%; RAM đỉnh 51,50 MiB"),
    ("CPU/RAM Server", "CPU TB 121,44%; RAM đỉnh 58,63 MiB", "CPU TB 231,55%; RAM đỉnh 64,05 MiB"),
    ("Lỗi", "0/10 file; 0%", "0/30 file; 0%"),
    ("Kết luận", "Thành công 10 file; không còn .part", "Thành công 30 file; không còn .part"),
]
add_table(doc, ["Chỉ số", "Mức tải 1", "Mức tải 2"], stress_rows, [4.1, 6.7, 6.7], font_size=8.8)
add_heading(doc, "8.1. Giới hạn của số liệu", 2)
add_body(doc, "Kết quả loopback phản ánh hiệu năng của protocol và máy đo, không đại diện trực tiếp cho LAN hoặc Internet. CPU vượt 100% nghĩa là tiến trình sử dụng nhiều lõi. Hai mức tải thay đổi đồng thời số Client và tổng dung lượng, vì vậy chưa thể tách riêng ảnh hưởng của từng biến.", size=11.5, line=1.15)
add_page_break(doc)

# Page 12 / numbered page 9
add_heading(doc, "9. KẾT QUẢ ĐẠT ĐƯỢC", 1)
add_heading(doc, "9.1. Luồng upload hoàn chỉnh", 2)
add_body(doc, "Phiên bản hiện tại thực hiện đầy đủ chuỗi xử lý: Client đọc thông tin file, tính hash, gửi metadata, chờ Ready, truyền dữ liệu theo chunk, nhận Completed và hiển thị tên Server đã lưu. Mỗi file có requestId và trạng thái riêng nên lỗi không làm dừng dispatcher.")
add_heading(doc, "9.2. Kết quả phía Client", 2)
for item in [
    "Chọn hoặc kéo-thả nhiều file; tự lập hàng đợi và giới hạn ba upload đồng thời.",
    "Hiển thị tiến độ, tốc độ, thông báo, tên file đã lưu và thống kê tổng hợp.",
    "Cancel và Retry từng file; thao tác hàng loạt không làm sai trạng thái của mục khác.",
    "Kiểm tra file thay đổi trong lúc tính hash hoặc trước khi gửi để dừng sớm dữ liệu không nhất quán.",
]:
    add_bullet(doc, item)
add_heading(doc, "9.3. Kết quả phía Server", 2)
for item in [
    "Kiểm tra version, requestId, tên file, kích thước, hash và trạng thái protocol.",
    "Áp dụng timeout cho metadata và dữ liệu; ghi log theo vòng đời Connect, Start, Completed, Error và Disconnect.",
    "Ghi file tạm, xác minh SHA-256, đổi tên khi trùng và không ghi đè file đã có.",
    "Dọn file .part trong các lỗi mà tiến trình còn hoạt động; shutdown có chờ session.",
]:
    add_bullet(doc, item)
add_heading(doc, "9.4. Kết quả kiểm chứng", 2)
add_body(doc, "Solution build Release với 0 warning, 0 error; bộ test kỹ thuật đạt 11/11. Stress test đạt 0 lỗi ở cả hai mức tải và không để lại file .part. Video demo ghi nhận thao tác GUI và kết quả lưu file; liên kết được trình bày tại mục 12.")
add_page_break(doc)

# Page 13 / numbered page 10
add_heading(doc, "10. KHÓ KHĂN VÀ KIẾN THỨC HỌC ĐƯỢC", 1)
add_heading(doc, "10.1. Khó khăn và cách xử lý", 2)
difficulty_rows = [
    ("Ranh giới dữ liệu trên TCP", "Dùng prefix 4 byte, hàm đọc đủ byte và fileSize cho phần raw binary"),
    ("GUI khi nhiều file chạy", "Tách queue, manager và ViewModel; dùng async/await, token riêng và SemaphoreSlim"),
    ("File thay đổi khi đang xử lý", "So sánh kích thước và LastWriteTimeUtc trước/sau hash và trước khi truyền"),
    ("WPF phụ thuộc Windows", "Tách Shared/Server để kiểm thử đa nền tảng; chạy GUI trên Windows hoặc Windows VM"),
    ("Phân biệt hủy và lỗi mạng", "CancellationToken quản lý thao tác người dùng; timeout và exception phản ánh lỗi truyền"),
]
add_table(doc, ["Khó khăn", "Cách xử lý"], difficulty_rows, [6.0, 11.5], font_size=10.3)
add_heading(doc, "10.2. Kiến thức và kỹ năng học được", 2)
for item in [
    "Hiểu TCP là luồng byte và ứng dụng phải tự định nghĩa framing, validation và vòng đời session.",
    "Thực hành async/await, CancellationToken, SemaphoreSlim, I/O theo chunk và cập nhật trạng thái WPF.",
    "Phân biệt kiểm tra toàn vẹn bằng SHA-256 với bảo mật bằng xác thực, phân quyền hoặc TLS.",
    "Thiết kế ca test có điều kiện đo, đọc log theo requestId và không suy rộng kết quả loopback sang LAN.",
    "Phối hợp module, tích hợp trên cùng solution và theo dõi tiến độ qua lịch sử commit.",
]:
    add_bullet(doc, item, size=12.2)
add_page_break(doc)

# Page 14 / numbered page 11
add_heading(doc, "11. NỘI DUNG ĐÃ VÀ CHƯA HOÀN THÀNH", 1)
add_caption(doc, "Bảng 11.1. Phạm vi triển khai của sản phẩm", kind="table")
scope_rows = [
    ("Đã hoàn thành trong code", "Kéo-thả nhiều file; queue; tối đa 3 upload/Client; tiến độ; Cancel/Retry; protocol V3; validation; timeout; SHA-256; .part; đổi tên trùng; log; thống kê"),
    ("Đã kiểm tra", "Build Release 0 warning, 0 error; 11/11 test kỹ thuật; stress 1 và 3 Client trên loopback; đối chiếu kích thước và SHA-256"),
    ("Cần nghiệm thu thêm", "LAN hoặc hai máy Windows; đo lại trong điều kiện mạng thật; kiểm tra quyền ghi và hết dung lượng trên nhiều cấu hình"),
    ("Chưa thực hiện", "Pause/Resume; upload thư mục; xác thực; TLS; mã hóa đầu cuối; lưu queue sau khi đóng Client; giới hạn tổng kết nối Server; tự dọn .part cũ sau khi tiến trình bị kill"),
]
add_table(doc, ["Trạng thái", "Nội dung"], scope_rows, [4.3, 13.2], font_size=10.2)
add_heading(doc, "11.1. Giới hạn hiện tại", 2)
add_body(doc, "Giới hạn ba file chỉ áp dụng cho một Client; Server chưa trả ServerBusy theo tổng tải. Retry gửi lại toàn bộ file, chưa resume từ offset. SaveDirectory là thư mục cục bộ. Khi tiến trình Server bị kill hoặc máy mất điện, code chưa quét file .part cũ lúc khởi động lại.")
add_heading(doc, "11.2. Hướng phát triển", 2)
for item in [
    "Bổ sung TLS, xác thực và phân quyền trước khi sử dụng ngoài môi trường học tập.",
    "Thêm resume theo offset, lưu hàng đợi và cơ chế dọn file .part khi khởi động.",
    "Áp dụng giới hạn tổng kết nối, backpressure và ServerBusy để bảo vệ Server.",
    "Đo LAN/Internet với nhiều cấu hình, lặp lại nhiều lần và báo cáo trung vị cùng P95.",
]:
    add_bullet(doc, item)
add_page_break(doc)

# Page 15 / numbered page 12
add_heading(doc, "12. KẾT LUẬN LIÊN KẾT VÀ TÀI LIỆU THAM KHẢO", 1)
add_heading(doc, "12.1. Kết luận", 2)
add_body(doc, "Đề tài đã hoàn thành mục tiêu cốt lõi: xây dựng ứng dụng upload nhiều file qua TCP Socket với giao diện kéo-thả, hàng đợi bất đồng bộ, trạng thái độc lập, Cancel/Retry, framing rõ ràng, kiểm tra SHA-256 và lưu file an toàn. Kết quả build, kiểm thử kỹ thuật và stress test cho thấy luồng chính hoạt động đúng trong điều kiện đã công bố.")
add_body(doc, "Sản phẩm vẫn là mô hình học tập, chưa phải dịch vụ lưu trữ sản xuất. Các giới hạn về TLS, xác thực, resume, backpressure và dọn file sau sự cố tiến trình đã được xác định để nhóm trình bày đúng phạm vi và tiếp tục phát triển.")
add_heading(doc, "12.2. Liên kết nộp bài", 2)
p = doc.add_paragraph()
format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.2, after=3, first_line=0)
r = p.add_run("Video demo: ")
set_run_font(r, size=13, bold=True)
add_hyperlink(p, "Mở video demo trên YouTube", "https://youtu.be/lFGhi14CTBE")
p = doc.add_paragraph()
format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.2, after=3, first_line=0)
r = p.add_run("GitHub repository: ")
set_run_font(r, size=13, bold=True)
add_hyperlink(p, "Mở repository trên GitHub", "https://github.com/lmn2005/UDM10_UploadFiles")
add_heading(doc, "12.3. Tài liệu tham khảo", 2)
references = [
    '[1] W. Eddy, Ed., "Transmission Control Protocol (TCP)," RFC 9293, Aug. 2022. [Online]. Available: https://www.rfc-editor.org/rfc/rfc9293.html',
    '[2] Microsoft, "NetworkStream Class," Microsoft Learn. [Online]. Available: https://learn.microsoft.com/dotnet/api/system.net.sockets.networkstream',
    '[3] National Institute of Standards and Technology, Secure Hash Standard, FIPS PUB 180-4, Aug. 2015. doi: 10.6028/NIST.FIPS.180-4.',
    '[4] Microsoft, "Asynchronous programming with async and await," Microsoft Learn. [Online]. Available: https://learn.microsoft.com/dotnet/csharp/asynchronous-programming/',
]
for ref in references:
    p = add_body(doc, ref, size=11.5, line=1.1, first_line=0, after=3)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.hanging_indent = Cm(0.75)
add_heading(doc, "12.4. Quy cách đóng gói", 2)
add_body(doc, "Khi nộp toàn bộ hồ sơ, nhóm xóa thư mục build và dependency cache không cần thiết, sau đó đóng gói theo tên 012012301304-Net3_Group_11-UDM_10.7z.", size=12, line=1.15, first_line=0)

# Set language and document compatibility options.
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

for p in doc.paragraphs:
    for run in p.runs:
        r_pr = run._element.get_or_add_rPr()
        lang = r_pr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            r_pr.append(lang)
        lang.set(qn("w:val"), "vi-VN")

doc.core_properties.title = "Báo cáo tiểu luận cuối kỳ ứng dụng upload nhiều file bằng kéo thả qua TCP Socket"
doc.core_properties.subject = "Lập trình mạng UDM 10"
doc.core_properties.author = "Net3 Group 11"
doc.core_properties.keywords = "TCP Socket, WPF, upload nhiều file, UDM 10"
doc.save(OUT)
print(OUT)
